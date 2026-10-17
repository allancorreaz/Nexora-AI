"""Extracao de texto dos formatos aceitos pela base de conhecimento Nexora AI.

Projeto criado por Allan de Melo Correa.
https://github.com/allancorreaz/Nexora-AI

Este modulo nao persiste arquivos. Upload, autorizacao de acesso e
limites de tamanho devem ser tratados antes da chamada a este parser.
"""

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx"})


class DocumentProcessingError(ValueError):
    """Falha de validacao ou extracao de um documento."""


def extract_text(filename: str, content: bytes) -> str:
    """Extrai texto sem gravar o documento no disco.

    O formato e selecionado pela extensao do nome enviado. A API que
    recebe o upload deve limitar o tamanho e validar os dados antes
    de chamar esta funcao. Documentos sem texto selecionavel (como
    PDFs digitalizados) precisam de OCR, ainda nao implementado.
    """
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentProcessingError(
            "Formato nao suportado. Utilize TXT, MD, PDF ou DOCX."
        )

    if not content:
        raise DocumentProcessingError("O documento esta vazio.")

    try:
        if extension in {".txt", ".md"}:
            extracted = content.decode("utf-8-sig")
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise DocumentProcessingError(
                    "PDF protegido por senha nao e suportado."
                )
            extracted = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            extracted = "\n".join(paragraphs)
    except DocumentProcessingError:
        raise
    except (UnicodeError, ValueError, OSError, KeyError, TypeError, RuntimeError) as exc:
        raise DocumentProcessingError(
            "Nao foi possivel ler o documento. Verifique o arquivo e o formato."
        ) from exc
    except Exception as exc:
        # Bibliotecas PDF/Office usam excecoes diferentes para arquivos corrompidos.
        raise DocumentProcessingError(
            "Nao foi possivel extrair o texto do documento."
        ) from exc

    extracted = extracted.strip()
    if not extracted:
        raise DocumentProcessingError(
            "Nenhum texto extraivel foi encontrado. PDFs digitalizados exigem OCR."
        )

    return extracted
