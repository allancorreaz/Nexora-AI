"""Testes do leitor de documentos da Nexora AI."""

import unittest
from io import BytesIO

from docx import Document
from pypdf import PdfWriter

from app.document_processor import DocumentProcessingError, extract_text


class DocumentProcessorTests(unittest.TestCase):
    def test_reads_utf8_text(self):
        self.assertEqual(
            extract_text("treinamento.txt", "Olá, treinamento!".encode("utf-8")),
            "Olá, treinamento!",
        )

    def test_reads_markdown_with_utf8_bom(self):
        data = b"\xef\xbb\xbf# Manual\n\nPrimeiro passo"
        self.assertEqual(
            extract_text("manual.MD", data),
            "# Manual\n\nPrimeiro passo",
        )

    def test_reads_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("Conferência de mercadorias")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Etapa"
        table.cell(0, 1).text = "Armazenamento"
        buffer = BytesIO()
        document.save(buffer)

        text = extract_text("processo.docx", buffer.getvalue())

        self.assertIn("Conferência de mercadorias", text)
        self.assertIn("Etapa | Armazenamento", text)

    def test_rejects_blank_pdf_without_extractable_text(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = BytesIO()
        writer.write(buffer)

        with self.assertRaisesRegex(DocumentProcessingError, "OCR"):
            extract_text("digitalizado.pdf", buffer.getvalue())

    def test_rejects_empty_file(self):
        with self.assertRaisesRegex(DocumentProcessingError, "vazio"):
            extract_text("arquivo.txt", b"")

    def test_rejects_unsupported_extension(self):
        with self.assertRaisesRegex(DocumentProcessingError, "nao suportado"):
            extract_text("arquivo.exe", b"conteudo")

    def test_rejects_invalid_utf8(self):
        with self.assertRaises(DocumentProcessingError):
            extract_text("arquivo.txt", b"\xff\xfe\xfd")


if __name__ == "__main__":
    unittest.main()
